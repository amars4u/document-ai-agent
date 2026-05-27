import os
from docx import Document
from docx.oxml.ns import qn

def get_outline_level(para):
    try:
        pPr = para._p.find(qn('w:pPr'))
        if pPr is not None:
            outlineLvl = pPr.find(qn('w:outlineLvl'))
            if outlineLvl is not None:
                return int(outlineLvl.get(qn('w:val')))
    except Exception:
        pass
    return None

def load_documents(folder="docs"):
    all_text = ""
    for filename in os.listdir(folder):
        if not filename.endswith(".docx"):
            continue
        path = os.path.join(folder, filename)
        try:
            doc = Document(path)
        except Exception as e:
            all_text += f"\n\n[Could not read {filename}: {e}]\n"
            continue

        all_text += f"\n\n{'='*50}\nDOCUMENT: {filename}\n{'='*50}\n"

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue
            level = get_outline_level(para)
            if level is not None:
                all_text += f"\n[SECTION: {text}]\n"
            else:
                if "\n" in para.text:
                    for line in para.text.split("\n"):
                        line = line.strip()
                        if line:
                            all_text += f"- {line}\n"
                else:
                    all_text += f"- {text}\n"

    return all_text


def load_and_chunk_documents(folder="docs", chunk_size=1000, chunk_overlap=200):
    chunks = []
    for filename in sorted(os.listdir(folder)):
        if not filename.endswith(".docx"):
            continue
        path = os.path.join(folder, filename)
        try:
            doc = Document(path)
        except Exception:
            continue

        current_section = "General"
        current_text = ""

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            level = get_outline_level(para)
            if level is not None:
                if current_text.strip():
                    chunks.append({
                        "text": current_text.strip(),
                        "source": filename,
                        "section": current_section,
                    })
                current_section = text
                current_text = ""
            else:
                current_text += text + "\n"
                if len(current_text) >= chunk_size:
                    chunks.append({
                        "text": current_text.strip(),
                        "source": filename,
                        "section": current_section,
                    })
                    current_text = current_text[-chunk_overlap:]

        if current_text.strip():
            chunks.append({
                "text": current_text.strip(),
                "source": filename,
                "section": current_section,
            })

    return chunks